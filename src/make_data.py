#!/usr/bin/env python3
"""최종 핸즈온용 더미 데이터 — 점포 진열 · 재고 점검 기록.

제품명은 공개된 농심 제품명을 쓰고, 점포 · 사람 · 수치는 모두 교육용 가상 데이터다.

출력
  docs/data/store_check_raw.xlsx    정제 전 원본 (표기 오류 · 중복 · 형식 섞임 · 이상값)
  docs/data/store_check_clean.xlsx  정답 정제본 (막혔을 때 대체용)
  docs/data/answer.json             정답 숫자

사용: python3 src/make_data.py   (openpyxl 필요 · pip install openpyxl)
같은 시드라 몇 번을 돌려도 같은 파일이 나온다.
"""
import json
import random
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "data"
R = random.Random(20261013)

STORES = [  # 코드, 점포명, 채널, 지역, 점검자
    ("S01", "A편의점 강남역점", "편의점", "서울", "김주현"), ("S02", "A편의점 판교점", "편의점", "경기 남부", "김주현"),
    ("S03", "B편의점 홍대점", "편의점", "서울", "김주현"), ("S04", "B편의점 수원역점", "편의점", "경기 남부", "김주현"),
    ("S05", "C편의점 잠실점", "편의점", "서울", "최민호"), ("S06", "C편의점 일산점", "편의점", "경기 북부", "최민호"),
    ("S07", "D편의점 부평점", "편의점", "인천", "최민호"), ("S08", "D편의점 의정부점", "편의점", "경기 북부", "최민호"),
    ("S09", "A마트 영등포점", "대형마트", "서울", "이지원"), ("S10", "A마트 동탄점", "대형마트", "경기 남부", "이지원"),
    ("S11", "B마트 송도점", "대형마트", "인천", "이지원"), ("S12", "B마트 파주점", "대형마트", "경기 북부", "이지원"),
    ("S13", "동부대리점", "대리점", "경기 남부", "박서연"), ("S14", "북부대리점", "대리점", "경기 북부", "박서연"),
]
PRODUCTS = [  # 코드, 제품명, 카테고리, 적정재고(박스)
    ("P01", "신라면", "봉지면", 40), ("P02", "짜파게티", "봉지면", 30), ("P03", "너구리", "봉지면", 24),
    ("P04", "안성탕면", "봉지면", 20), ("P05", "신라면 블랙", "봉지면", 16), ("P06", "육개장사발면", "컵면", 30),
    ("P07", "신라면 큰사발", "컵면", 24), ("P08", "새우깡", "스낵", 30), ("P09", "양파링", "스낵", 20),
    ("P10", "꿀꽈배기", "스낵", 18),
]
CH_SCALE = {"편의점": 1.0, "대형마트": 2.5, "대리점": 4.0}
DATES = [date(2026, 7, 6) + timedelta(weeks=w) for w in range(13)]  # 7/6 ~ 9/28, 주 1회

STORE_TYPO = {
    "A편의점 강남역점": ["A편의점 강남점", "A 편의점 강남역점"], "A편의점 판교점": ["A편의점판교점", "A편의점 판교 점"],
    "B편의점 홍대점": ["B편의졈 홍대점", "B편의점 홍대입구점"], "B편의점 수원역점": ["B편의점 수원점"],
    "C편의점 잠실점": ["C편의점  잠실점"], "C편의점 일산점": ["C 편의점 일산점"], "D편의점 부평점": ["D편의점 부평역점"],
    "D편의점 의정부점": ["D편의점 의정부"], "A마트 영등포점": ["A마트영등포점", "A 마트 영등포점"],
    "A마트 동탄점": ["A마트 동탄"], "B마트 송도점": ["B마트(송도)"], "B마트 파주점": ["B마트 파주 점"],
    "동부대리점": ["동부 대리점"], "북부대리점": ["북부대리점점"],
}
PRODUCT_TYPO = {"신라면": ["신 라면", "辛라면"], "짜파게티": ["짜파게뤼", "짜파게티 "], "신라면 블랙": ["신라면블랙", "신라면 BLACK"],
                "육개장사발면": ["육개장 사발면"], "신라면 큰사발": ["신라면큰사발", "신라면 큰사발면"], "새우깡": ["새우 깡"],
                "양파링": ["양파 링"], "꿀꽈배기": ["꿀 꽈배기"], "너구리": ["너구리 "], "안성탕면": ["안성 탕면"]}
REGION_TYPO = {"서울": ["서울시", "서울특별시"], "경기 남부": ["경기남부", "경기(남)"], "경기 북부": ["경기북부"], "인천": ["인천시"]}
YN_TYPO = {"Y": ["O", "예", "있음", "y"], "N": ["X", "아니오", "없음", "n", ""]}


def clean_rows():
    rows = []
    for d in DATES:
        wk = (d - DATES[0]).days // 7
        promo_week = 4 <= wk <= 7                                  # 8월 경쟁사 스낵 2+1 행사 기간
        for code, store, ch, region, who in STORES:
            for pc, pname, cat, need in PRODUCTS:
                if ch == "대리점" and cat == "컵면":
                    continue
                target = int(round(need * CH_SCALE[ch]))
                stock = target * R.uniform(0.55, 1.35)
                if pname == "신라면" and region == "경기 남부" and 4 <= wk <= 6:
                    stock *= 0.25                                   # 8월 물류 지연 → 결품
                if pname == "신라면 블랙" and wk >= 9:
                    stock *= 1.4                                    # 9월 행사 대비 입고
                comp = "Y" if (cat == "스낵" and promo_week and ch == "편의점" and R.random() < 0.7) else "N"
                face = max(1, int(round((4 if cat != "스낵" else 3) * R.uniform(0.8, 1.2) * (0.6 if comp == "Y" else 1))))
                if comp == "Y":
                    stock *= 0.7
                stock = int(round(stock))
                if stock < target * 0.3 or R.random() < 0.03:
                    stock = R.choice([0, 0, 1, 2])
                if pname == "신라면" and region == "경기 남부" and 4 <= wk <= 6 and R.random() < 0.8:
                    stock = 0
                if pname == "육개장사발면" and ch == "편의점" and wk in (5, 6) and R.random() < 0.35:
                    stock = 0                                       # 여름 휴가철 수요
                promo = "Y" if (pname == "신라면 블랙" and wk >= 10) else "N"
                rows.append({"점검일": d.isoformat(), "지역": region, "점포코드": code, "점포명": store, "채널": ch,
                             "제품코드": pc, "제품명": pname, "카테고리": cat, "진열면수": face, "적정재고": target,
                             "실재고": stock, "결품": "Y" if stock == 0 else "N", "자사행사": promo, "경쟁사행사": comp,
                             "점검자": who})
    return rows


def dirty(rows):
    log = defaultdict(int)
    out = []
    for r in rows:
        d = dict(r)
        if R.random() < 0.06:
            d["점포명"] = R.choice(STORE_TYPO[r["점포명"]]); log["점포명 표기"] += 1
        if R.random() < 0.05:
            d["제품명"] = R.choice(PRODUCT_TYPO[r["제품명"]]); log["제품명 표기"] += 1
        if R.random() < 0.05:
            d["지역"] = R.choice(REGION_TYPO[r["지역"]]); log["지역 표기"] += 1
        if R.random() < 0.08:
            y, m, dd = r["점검일"].split("-")
            d["점검일"] = R.choice([f"{y}.{m}.{dd}", f"{int(m)}/{int(dd)}", f"{y[2:]}년 {int(m)}월 {int(dd)}일",
                                    date(int(y), int(m), int(dd))]); log["날짜 형식"] += 1
        if R.random() < 0.10:
            d["결품"] = R.choice(YN_TYPO[r["결품"]]); log["Y/N 표기"] += 1
        if R.random() < 0.05:
            d["실재고"] = f'{r["실재고"]}박스'; log["숫자가 문자"] += 1
        if R.random() < 0.008:
            d["실재고"] = r["실재고"] * 10 + R.randint(0, 9); log["이상값(자릿수 실수)"] += 1
        if R.random() < 0.01:
            d["점포코드"] = ""; log["점포코드 빈칸"] += 1
        out.append(d)
    for r in R.sample(out, 30):                                    # 같은 점검을 두 번 입력
        out.insert(R.randrange(len(out)), dict(r)); log["중복 점검"] += 1
    for _ in range(8):
        out.insert(R.randrange(len(out)), None); log["빈 행"] += 1
    return out, dict(log)


COLS = ["점검일", "지역", "점포코드", "점포명", "채널", "제품코드", "제품명", "카테고리",
        "진열면수", "적정재고", "실재고", "결품", "자사행사", "경쟁사행사", "점검자"]
NOTICE = "교육용 가상 데이터: 제품명은 공개 제품명이며, 점포 · 사람 · 수치는 모두 가상입니다."


def style_header(ws, row):
    for c in ws[row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(horizontal="center")


def add_common(wb):
    ws = wb.create_sheet("점포_마스터")
    ws.append(["점포코드", "점포명", "채널", "지역", "점검자"]); style_header(ws, 1)
    for s in STORES:
        ws.append(list(s))
    ws = wb.create_sheet("제품_마스터")
    ws.append(["제품코드", "제품명", "카테고리", "적정재고(편의점 기준)"]); style_header(ws, 1)
    for p in PRODUCTS:
        ws.append(list(p))
    ws = wb.create_sheet("읽어두기")
    for line in [NOTICE, "기간: 2026-07-06 ~ 2026-09-28 · 주 1회 점검 · 점포 14곳 · 제품 10종",
                 "적정재고 = 제품 기준 × 채널 배수 (편의점 1 · 대형마트 2.5 · 대리점 4)",
                 "결품 = 실재고 0 · 재고충족률 = 실재고 ÷ 적정재고",
                 "경쟁사행사 = 점검 당일 같은 매대에 경쟁사 행사가 있었는지"]:
        ws.append([line])


def answer(rows, raw, log):
    def rate(sel):
        sel = list(sel)
        return round(sum(r["결품"] == "Y" for r in sel) / len(sel) * 100, 1) if sel else 0
    aug = [r for r in rows if r["점검일"][:7] == "2026-08"]
    by_prod = {p[1]: rate(r for r in aug if r["제품명"] == p[1]) for p in PRODUCTS}
    by_region = {g: rate(r for r in aug if r["지역"] == g) for g in ["서울", "경기 남부", "경기 북부", "인천"]}
    snack = [r for r in rows if r["카테고리"] == "스낵" and r["채널"] == "편의점"]
    face = lambda c: round(sum(r["진열면수"] for r in snack if r["경쟁사행사"] == c) / max(1, sum(r["경쟁사행사"] == c for r in snack)), 2)
    fill = round(sum(min(r["실재고"] / r["적정재고"], 2) for r in rows) / len(rows) * 100, 1)
    worst = max(by_prod, key=by_prod.get)
    return {
        "원본 데이터 행 수(빈 행 제외)": len(raw),
        "정제 후 행 수": len(rows),
        "섞은 오류": log,
        "전체 결품률(%)": rate(rows),
        "8월 결품률(%)": rate(aug),
        "8월 결품률 1위 제품": f"{worst} {by_prod[worst]}%",
        "8월 지역별 결품률(%)": by_region,
        "편의점 스낵 평균 진열면수(경쟁사 행사 有/無)": [face("Y"), face("N")],
        "평균 재고충족률(%)": fill,
    }


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows = clean_rows()
    raw, log = dirty(rows)

    wb = Workbook()
    ws = wb.active; ws.title = "점검_원본"
    ws.append(["2026년 3분기 점포 진열 · 재고 점검 기록 (수도권영업팀)"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([NOTICE])
    ws.append(COLS); style_header(ws, 3)
    for r in raw:
        ws.append([] if r is None else [r[c] for c in COLS])
    add_common(wb)
    wb.save(OUT / "store_check_raw.xlsx")

    wb = Workbook()
    ws = wb.active; ws.title = "점검_정제"
    ws.append(COLS); style_header(ws, 1)
    for r in rows:
        ws.append([r[c] for c in COLS])
    add_common(wb)
    wb.save(OUT / "store_check_clean.xlsx")

    ans = answer(rows, [r for r in raw if r is not None], log)
    (OUT / "answer.json").write_text(json.dumps(ans, ensure_ascii=False, indent=2))
    print(json.dumps(ans, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
