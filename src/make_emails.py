#!/usr/bin/env python3
"""메일 Skill 실습용 「내가 보낸 메일」 예시 5통. 사람 · 점포 · 일정은 모두 가상이다.

출력
  docs/data/my_emails_sample.txt   Skill 참고 파일로 바로 올리는 텍스트
  docs/data/my_emails_sample.docx  Drive · Docs로 열어 고쳐 쓰는 문서

사용: python3 src/make_emails.py   (docx는 python-docx 필요 · 없으면 txt만 만든다)
"""
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "docs" / "data"

HEAD = ("내가 보낸 메일 예시 (교육용 가상 자료)\n"
        "메일 Skill의 참고 파일로 씁니다. 내 실제 메일로 바꿔 넣을 때는 이름 · 연락처 · 거래 조건을 가립니다.\n")

SIGN = "감사합니다.\n김주현 드림\n수도권영업팀 · 편의점 담당"

MAILS = [
    ("미팅 후 감사 · 후속 일정",
     "C편의점 담당자D",
     "8/26 미팅 감사 및 후속 일정 안내",
     "담당자님, 안녕하세요. 수도권영업팀 김주현입니다.\n\n"
     "어제 미팅에 시간 내주셔서 감사합니다.\n"
     "말씀하신 내용은 아래와 같이 정리했습니다.\n\n"
     "- 육개장사발면 진열 확대: 9월 미팅에서 점포 수 · 위치 확정\n"
     "- 신라면 블랙 가을 진열 제안서: 9/12(금)까지 전달\n\n"
     "빠진 내용이 있으면 편하게 말씀 주세요."),
    ("결품 문의 1차 회신",
     "A편의점 담당자C",
     "[회신] 신라면 결품 문의 건",
     "담당자님, 안녕하세요. 수도권영업팀 김주현입니다.\n\n"
     "문의 주신 신라면 결품 건 확인했습니다.\n"
     "현재 물류팀에 입고 일정과 납품 내역 확인을 요청해 두었습니다.\n\n"
     "확인 결과를 받는 대로 2영업일 안에 다시 말씀드리겠습니다.\n"
     "그 전에 필요한 자료가 있으면 알려 주세요.\n\n"
     "불편을 드려 죄송합니다."),
    ("행사 제안서 송부",
     "B편의점 담당자A",
     "9월 새우깡 행사 제안서 송부",
     "담당자님, 안녕하세요. 수도권영업팀 김주현입니다.\n\n"
     "지난 미팅에서 말씀드린 9월 새우깡 행사 제안서를 첨부드립니다.\n\n"
     "- 행사: 2+1 · 9/15~10/12 (4주)\n"
     "- 대상: 홍대점 · 수원역점 등 수도권 점포\n"
     "- 조건: 행사 발주 300박스 이상\n\n"
     "검토하시고 의견 주시면 일정에 맞춰 준비하겠습니다."),
    ("미팅 일정 조율",
     "D편의점 담당자E",
     "9월 정기 미팅 일정 문의",
     "담당자님, 안녕하세요. 수도권영업팀 김주현입니다.\n\n"
     "9월 정기 미팅 일정을 여쭙고자 연락드립니다.\n"
     "아래 시간 중 편하신 때를 알려 주시면 맞추겠습니다.\n\n"
     "- 9/9(화) 오후 2시\n"
     "- 9/10(수) 오전 10시\n"
     "- 9/11(목) 오후 4시\n\n"
     "다른 시간이 편하시면 그에 맞추겠습니다."),
    ("사내 자료 요청",
     "물류팀 담당자",
     "[요청] 경기 남부 신라면 입고 일정 확인",
     "안녕하세요, 수도권영업팀 김주현입니다.\n\n"
     "경기 남부 편의점에서 8월 둘째 주부터 신라면 결품이 이어지고 있습니다.\n"
     "아래 두 가지를 확인 부탁드립니다.\n\n"
     "- 8월 경기 남부 물류센터 신라면 입고 일정\n"
     "- 판교점 · 수원역점 8월 납품 내역\n\n"
     "거래처 회신 기한이 9/3(수)이라 9/2(화)까지 받을 수 있으면 좋겠습니다."),
]

STYLE = ("\n[내 메일의 공통점]\n"
         "- 호칭: 거래처는 「담당자님」, 사내는 생략\n"
         "- 첫인사: 「안녕하세요. 수도권영업팀 김주현입니다.」\n"
         "- 용건을 첫 문단에, 한 단락은 두 문장 이내\n"
         "- 일정 · 조건 · 요청 항목은 「- 」 목록으로\n"
         "- 확정되지 않은 일은 「~하겠습니다」 대신 「확인 후 말씀드리겠습니다」\n"
         "- 맺음: 「감사합니다. / 김주현 드림 / 수도권영업팀 · 편의점 담당」\n")


def text():
    parts = [HEAD]
    for i, (kind, to, subj, body) in enumerate(MAILS, 1):
        parts.append(f"\n========== 예시 {i} · {kind} ==========\n받는 사람: {to}\n제목: {subj}\n\n{body}\n\n{SIGN}\n")
    parts.append(STYLE)
    return "".join(parts)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "my_emails_sample.txt").write_text(text(), encoding="utf-8")
    try:
        from docx import Document
    except ImportError:
        print("python-docx 없음 → txt만 만듦")
        return
    doc = Document()
    doc.add_heading("내가 보낸 메일 예시 (교육용 가상 자료)", 1)
    doc.add_paragraph(HEAD.split("\n")[1])
    for i, (kind, to, subj, body) in enumerate(MAILS, 1):
        doc.add_heading(f"예시 {i} · {kind}", 2)
        doc.add_paragraph(f"받는 사람: {to}\n제목: {subj}")
        for para in (body + "\n\n" + SIGN).split("\n\n"):
            doc.add_paragraph(para)
    doc.add_heading("내 메일의 공통점", 2)
    for line in STYLE.strip().split("\n")[1:]:
        doc.add_paragraph(line.lstrip("- "), style="List Bullet")
    doc.save(OUT / "my_emails_sample.docx")
    print(OUT / "my_emails_sample.txt", OUT / "my_emails_sample.docx")


if __name__ == "__main__":
    main()
